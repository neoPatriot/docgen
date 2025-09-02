import os
from flask import Flask, render_template, request, send_from_directory, flash, redirect, url_for, session
from werkzeug.utils import secure_filename
import openpyxl
from docx import Document
import zipfile
import uuid

app = Flask(__name__)
app.config['UPLOAD_FOLDER'] = 'uploads'
app.config['SECRET_KEY'] = 'supersecretkey'
os.makedirs(app.config['UPLOAD_FOLDER'], exist_ok=True)

@app.route('/')
def index():
    download_file = session.pop('download_file', None)
    return render_template('index.html', download_file=download_file)

def generate_documents(template_path, data_path, output_folder):
    """
    Generates DOCX files from a template and Excel data.
    Returns the number of documents generated.
    """
    try:
        workbook = openpyxl.load_workbook(data_path)
        sheet = workbook.active

        header = [cell.value for cell in sheet[1]]

        doc_count = 0
        for i, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            doc = Document(template_path)
            row_data = dict(zip(header, row))

            for key, value in row_data.items():
                placeholder = f"{{{{{key}}}}}"
                str_value = str(value) if value is not None else ""

                for para in doc.paragraphs:
                    if placeholder in para.text:
                        # Using runs to preserve formatting
                        inline = para.runs
                        for j in range(len(inline)):
                            if placeholder in inline[j].text:
                                text = inline[j].text.replace(placeholder, str_value)
                                inline[j].text = text

                for table in doc.tables:
                    for t_row in table.rows:
                        for cell in t_row.cells:
                            for para in cell.paragraphs:
                                if placeholder in para.text:
                                    inline = para.runs
                                    for j in range(len(inline)):
                                        if placeholder in inline[j].text:
                                            text = inline[j].text.replace(placeholder, str_value)
                                            inline[j].text = text

            # Generate a unique filename for each document
            output_filename = f"document_{i-1}.docx"
            output_path = os.path.join(output_folder, output_filename)
            doc.save(output_path)
            doc_count += 1

        return doc_count
    except Exception as e:
        print(f"An error occurred: {e}")
        return -1

@app.route('/upload', methods=['POST'])
def upload_files():
    if 'template_file' not in request.files or 'data_file' not in request.files:
        flash('Не удалось загрузить файлы.', 'error')
        return redirect(url_for('index'))

    template_file = request.files['template_file']
    data_file = request.files['data_file']

    if template_file.filename == '' or data_file.filename == '':
        flash('Файлы не выбраны.', 'error')
        return redirect(url_for('index'))

    if template_file and data_file and template_file.filename.endswith('.docx') and data_file.filename.endswith('.xlsx'):
        template_filename = secure_filename(template_file.filename)
        data_filename = secure_filename(data_file.filename)

        template_path = os.path.join(app.config['UPLOAD_FOLDER'], template_filename)
        data_path = os.path.join(app.config['UPLOAD_FOLDER'], data_filename)

        template_file.save(template_path)
        data_file.save(data_path)

        # Create a temporary directory for generated docs
        output_folder = os.path.join(app.config['UPLOAD_FOLDER'], 'generated_docs')
        os.makedirs(output_folder, exist_ok=True)

        # Generate documents
        num_docs = generate_documents(template_path, data_path, output_folder)

        if num_docs > 0:
            # Create a zip file with a unique name
            zip_filename = f"generated_documents_{uuid.uuid4().hex}.zip"
            zip_path = os.path.join(app.config['UPLOAD_FOLDER'], zip_filename)
            with zipfile.ZipFile(zip_path, 'w') as zipf:
                for root, _, files in os.walk(output_folder):
                    for file in files:
                        zipf.write(os.path.join(root, file), file)

            # Clean up temporary files
            os.remove(template_path)
            os.remove(data_path)
            for file in os.listdir(output_folder):
                os.remove(os.path.join(output_folder, file))
            os.rmdir(output_folder)

            flash(f'Успешно сгенерировано {num_docs} документов.', 'success')
            session['download_file'] = zip_filename
            return redirect(url_for('index'))
        else:
            flash('Произошла ошибка при генерации документов.', 'error')
            # Clean up uploaded files on error
            os.remove(template_path)
            os.remove(data_path)
            return redirect(url_for('index'))

    else:
        flash('Неверный формат файлов. Пожалуйста, загрузите .docx и .xlsx файлы.', 'error')
        return redirect(url_for('index'))

@app.route('/download/<filename>')
def download_file(filename):
    return send_from_directory(app.config['UPLOAD_FOLDER'], filename, as_attachment=True)


if __name__ == '__main__':
    app.run(debug=True)
